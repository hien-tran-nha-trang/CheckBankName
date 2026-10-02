"""Xử lý bảng Excel: chuẩn hoá STK, nhận diện ngân hàng, tra cứu tên, so khớp."""
from __future__ import annotations

import difflib
import re
import time
from typing import Callable, Iterable

import pandas as pd

from .banks import BankMatcher, normalize, unsupported_reason
from .lookup import LookupResult

ACCOUNT_HINTS = ["so tai khoan", "stk", "so tk", "tai khoan", "account", "acc no", "account number"]
BANK_HINTS = ["ngan hang", "ten ngan hang", "bank", "nh"]
NAME_HINTS = ["ten chu tai khoan", "chu tai khoan", "ten tai khoan", "ho ten", "ho va ten", "ten", "name"]

COL_BANK_FOUND = "NH nhận diện"
COL_RESULT_NAME = "Tên tra cứu"
COL_STATUS = "Trạng thái"
COL_MATCH = "So khớp tên"
RESULT_COLUMNS = [COL_BANK_FOUND, COL_RESULT_NAME, COL_STATUS, COL_MATCH]


def clean_account(value) -> str:
    """'0123 456.789' -> '0123456789'; 1.23e+11 (số Excel) -> '123000000000'."""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    if re.fullmatch(r"\d+\.0", text):
        text = text[:-2]
    return re.sub(r"[\s.\-_,']", "", text)


def guess_column(columns: Iterable[str], hints: list[str]) -> str | None:
    cols = list(columns)
    normed = {c: normalize(c) for c in cols}
    for hint in hints:
        for c, n in normed.items():
            if n == hint:
                return c
    for hint in hints:
        for c, n in normed.items():
            if re.search(rf"\b{re.escape(hint)}\b", n):
                return c
    return None


def guess_name_columns(columns: Iterable[str], exclude: Iterable[str] = ()) -> list[str]:
    """Các cột chứa tên người/đơn vị: 'Tên KH', 'Tên Đơn Vị', 'Tên Liên Hệ', 'Họ tên'..."""
    skip = set(exclude)
    found = []
    for c in columns:
        n = normalize(c)
        if c in skip or "ngan hang" in n or "bank" in n:
            continue
        if n.startswith("ten") or "ho ten" in n or "name" in n or "chu tai khoan" in n:
            found.append(c)
    return found


def compare_names(expected: str, actual: str) -> str:
    if not expected or not actual:
        return ""
    a, b = normalize(expected), normalize(actual)
    if a == b:
        return "Khớp"
    ratio = difflib.SequenceMatcher(None, a, b).ratio()
    if ratio >= 0.85:
        return f"Gần khớp ({ratio:.0%})"
    return f"KHÔNG khớp ({ratio:.0%})"


def best_match(row, name_cols: list[str], actual: str) -> str:
    """So tên tra cứu với nhiều cột (Tên KH, Tên đơn vị, Tên liên hệ...), lấy kết quả tốt nhất."""
    best, best_score = "", -1.0
    for col in name_cols:
        expected = row.get(col)
        if expected is None or pd.isna(expected) or not str(expected).strip():
            continue
        verdict = compare_names(str(expected), actual)
        score = 2.0 if verdict == "Khớp" else difflib.SequenceMatcher(
            None, normalize(expected), normalize(actual)).ratio()
        if score > best_score:
            best, best_score = (f"{verdict} - {col}" if len(name_cols) > 1 else verdict), score
    return best


def process(
    df: pd.DataFrame,
    account_col: str,
    bank_col: str,
    lookup_fn: Callable[[str, str], LookupResult],
    matcher: BankMatcher,
    name_col: str | list[str] | None = None,
    delay: float = 0.5,
    progress: Callable[[int, int], None] | None = None,
) -> pd.DataFrame:
    name_cols = [name_col] if isinstance(name_col, str) else list(name_col or [])
    out = df.copy()
    for c in RESULT_COLUMNS:
        out[c] = ""
    cache: dict[tuple[str, str], LookupResult] = {}
    total = len(out)
    for i, (idx, row) in enumerate(out.iterrows(), start=1):
        account = clean_account(row.get(account_col))
        bank_text = row.get(bank_col)
        bank = matcher.match(bank_text) if not pd.isna(bank_text) else None

        if bank:
            out.at[idx, COL_BANK_FOUND] = bank.label
        if not account:
            out.at[idx, COL_STATUS] = "Thiếu số tài khoản"
        elif not account.isdigit() or not 6 <= len(account) <= 19:
            out.at[idx, COL_STATUS] = "STK không hợp lệ (chỉ gồm 6-19 chữ số)"
        elif not bank:
            reason = unsupported_reason(bank_text) if not pd.isna(bank_text) else None
            out.at[idx, COL_STATUS] = reason or f"Không nhận diện được ngân hàng: '{bank_text}'"
        elif not bank.lookup_supported:
            out.at[idx, COL_STATUS] = f"{bank.short_name} chưa hỗ trợ tra cứu"
        else:
            key = (bank.bin, account)
            if key not in cache:
                cache[key] = lookup_fn(bank.bin, account)
                if delay:
                    time.sleep(delay)
            res = cache[key]
            out.at[idx, COL_RESULT_NAME] = res.account_name
            out.at[idx, COL_STATUS] = "OK" if res.ok else res.message
            if name_cols and res.ok:
                out.at[idx, COL_MATCH] = best_match(row, name_cols, res.account_name)
        if progress:
            progress(i, total)
    return out


def write_to_workbook(src, result: pd.DataFrame, header: str = "Tên TK tra cứu",
                      header_row: int = 1, sheet_name=0) -> bytes:
    """Ghi tên tra cứu vào file Excel gốc (giữ nguyên định dạng), thêm 1 cột ở cuối.

    Dòng lỗi để trống, tô đỏ và ghi lý do vào chú thích (comment) của ô.
    Dòng có tên nhưng không khớp với tên trong file được tô vàng.
    `result` phải được đọc từ chính file/sheet này (index = số dòng dữ liệu, bắt đầu từ 0).
    """
    import io

    import openpyxl
    from openpyxl.comments import Comment
    from openpyxl.styles import Font, PatternFill

    wb = openpyxl.load_workbook(src)
    ws = wb.worksheets[sheet_name] if isinstance(sheet_name, int) else wb[sheet_name]
    col = ws.max_column + 1
    head = ws.cell(row=header_row, column=col, value=header)
    prev = ws.cell(row=header_row, column=col - 1)
    if prev.has_style:
        head._style = prev._style
    head.font = Font(bold=True)
    red = PatternFill("solid", fgColor="FFC7CE")
    yellow = PatternFill("solid", fgColor="FFEB9C")
    for idx, row in result.iterrows():
        cell = ws.cell(row=header_row + 1 + int(idx), column=col)
        status = row.get(COL_STATUS, "")
        if status == "OK":
            cell.value = row[COL_RESULT_NAME]
            match = str(row.get(COL_MATCH, "") or "")
            if match and not match.startswith("Khớp"):
                cell.fill = yellow
                cell.comment = Comment(match, "CheckBankName")
        else:
            cell.fill = red
            cell.comment = Comment(str(status), "CheckBankName")
    ws.column_dimensions[head.column_letter].width = 40
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def read_excel(file, sheet_name=0) -> pd.DataFrame:
    """Đọc mọi cột dạng chuỗi để không mất số 0 đầu của STK."""
    return pd.read_excel(file, sheet_name=sheet_name, dtype=str)


def to_excel_bytes(df: pd.DataFrame) -> bytes:
    import io

    from openpyxl.styles import PatternFill

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Ket qua")
        ws = writer.sheets["Ket qua"]
        status_idx = list(df.columns).index(COL_STATUS) + 1
        green = PatternFill("solid", fgColor="C6EFCE")
        red = PatternFill("solid", fgColor="FFC7CE")
        for r in range(2, len(df) + 2):
            cell = ws.cell(row=r, column=status_idx)
            cell.fill = green if cell.value == "OK" else red
        for col in ws.columns:
            width = max(len(str(c.value or "")) for c in col) + 2
            ws.column_dimensions[col[0].column_letter].width = min(width, 50)
    return buf.getvalue()
