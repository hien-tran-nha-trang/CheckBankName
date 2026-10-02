import pandas as pd
import pytest

from bank_checker import BankMatcher, LookupResult, load_banks, process
from bank_checker.processor import clean_account, compare_names


@pytest.fixture(scope="module")
def matcher():
    return BankMatcher(load_banks(online=False))


@pytest.mark.parametrize("text,code", [
    ("Vietcombank", "VCB"), ("VCB", "VCB"), ("NH TMCP Ngoại thương Việt Nam", "VCB"),
    ("Ngân hàng Vietcombank - CN Nha Trang", "VCB"), ("970436", "VCB"),
    ("Vietinbank", "ICB"), ("Agribank", "VBA"), ("BIDV", "BIDV"), ("MB Bank", "MB"),
    ("mbbank", "MB"), ("Techcombank", "TCB"), ("ACB", "ACB"), ("Sacombank", "STB"),
    ("TP Bank", "TPB"), ("LienVietPostBank", "LPB"), ("Quân đội", "MB"),
    ("Ngân hàng Á Châu", "ACB"), ("SCB", "SCB"), ("Shinhan", "SHBVN"), ("SHB", "SHB"),
    ("Techcombnk", "TCB"),
])
def test_match(matcher, text, code):
    assert matcher.match(text).code == code


def test_match_unknown(matcher):
    assert matcher.match("abc xyz") is None
    assert matcher.match("") is None


def test_clean_account():
    assert clean_account(" 0123 456.789 ") == "0123456789"
    assert clean_account(123456789.0) == "123456789"
    assert clean_account("123456789.0") == "123456789"
    assert clean_account(None) == ""


def test_compare_names():
    assert compare_names("Nguyễn Văn A", "NGUYEN VAN A") == "Khớp"
    assert compare_names("Nguyen Van An", "NGUYEN VAN AN.").startswith("Khớp")
    assert compare_names("Tran Thi B", "NGUYEN VAN A").startswith("KHÔNG")


def test_process(matcher):
    df = pd.DataFrame({
        "STK": ["0071000123456", "abc", "123456789", "", "0071000123456"],
        "Ngân hàng": ["VCB", "ACB", "Ngân hàng lạ", "BIDV", "Vietcombank"],
        "Họ tên": ["Nguyễn Văn A", "", "", "", "Tran Van B"],
    })
    calls = []

    def fake(bin_code, acc):
        calls.append((bin_code, acc))
        return LookupResult(True, account_name="NGUYEN VAN A")

    out = process(df, "STK", "Ngân hàng", fake, matcher, name_col="Họ tên", delay=0)
    assert calls == [("970436", "0071000123456")]  # có cache, chỉ gọi 1 lần
    assert list(out["Trạng thái"]) [0] == "OK"
    assert out["So khớp tên"][0] == "Khớp"
    assert "không hợp lệ" in out["Trạng thái"][1]
    assert "Không nhận diện" in out["Trạng thái"][2]
    assert out["Trạng thái"][3] == "Thiếu số tài khoản"
    assert out["So khớp tên"][4].startswith("KHÔNG")


@pytest.mark.parametrize("text,code", [
    ("Ngân Hàng TMCP Á Châu_CN Sài Gòn", "ACB"),
    ("Ngân Hàng TMCP Phát Triển Nhà Thành Phố Hồ Chí Minh - CN Nha Trang - PGD Cam Ranh", "HDB"),
    ("Ngân Hàng TMCP Sài Gòn- Hà Nội_CN Khánh Hòa_PGD Cam Ranh", "SHB"),
    ("Ngân Hàng TMCP Sài Gòn Thương Tín_CN Khánh Hoà", "STB"),
    ("Ngân Hàng TMCP Đông Nam Á_CN Nha Trang", "SEAB"),
    ("Ngân Hàng Liên Doanh Việt - Nga_CN Khánh Hòa", "VRB"),
    ("Ngân Hàng\xa0TMCP\xa0Á Châu_CN Khánh Hòa", "ACB"),
    ("TECHCOMBANK_SÀI GÒN", "TCB"),
    ("Ngân hàng TMCP Việt Nam Thương Tín_CN Khánh Hoà", "VIETBANK"),
    ("Ngân hàng TMCP Phương Đông - Chi nhánh Phú Yên", "OCB"),
])
def test_match_with_branch(matcher, text, code):
    assert matcher.match(text).code == code


def test_kho_bac_not_supported(matcher):
    df = pd.DataFrame({"STK": ["3712123456"], "NH": ["Kho Bạc Nhà Nước Phú Yên"]})
    out = process(df, "STK", "NH", lambda b, a: LookupResult(True, "X"), matcher, delay=0)
    assert out["Trạng thái"][0].startswith("Kho bạc")


def test_write_to_workbook(tmp_path, matcher):
    import io

    import openpyxl

    from bank_checker.processor import guess_name_columns, write_to_workbook

    src = tmp_path / "in.xlsx"
    pd.DataFrame({
        "Tên KH": ["Nguyễn Văn A", "Lê B", "C"],
        "Tên Ngân Hàng": ["Vietcombank_CN Nha Trang", "ACB", "Kho Bạc Nhà Nước"],
        "Số TK Ngân Hàng": ["0071000123456", "12345678", "999999"],
    }).to_excel(src, index=False)
    df = pd.read_excel(src, dtype=str)
    assert guess_name_columns(df.columns, ["Tên Ngân Hàng"]) == ["Tên KH"]
    out = process(df, "Số TK Ngân Hàng", "Tên Ngân Hàng",
                  lambda b, a: LookupResult(True, "NGUYEN VAN A"), matcher, name_col=["Tên KH"], delay=0)
    ws = openpyxl.load_workbook(io.BytesIO(write_to_workbook(src, out))).active
    assert ws["D1"].value == "Tên TK tra cứu"
    assert ws["D2"].value == "NGUYEN VAN A" and ws["D2"].comment is None
    assert ws["D3"].value == "NGUYEN VAN A" and "KHÔNG khớp" in ws["D3"].comment.text
    assert ws["D4"].value is None and "Kho bạc" in ws["D4"].comment.text


def test_tracuubank_lookup(monkeypatch):
    from bank_checker import TracuuBankLookup, load_banks

    sent = {}

    class Resp:
        status_code = 200

        def __init__(self, body):
            self.body = body

        def json(self):
            return self.body

    def fake_get(url, params, timeout):
        sent.update(params)
        if params["bank_number"] == "1":
            return Resp({"status": "error", "message": "Không tìm thấy tài khoản"})
        return Resp({"status": "success", "data": {"accountName": "NGUYEN VAN A ", "bankCode": "VCB"}})

    lk = TracuuBankLookup("key", load_banks(online=False))
    monkeypatch.setattr(lk.session, "get", fake_get)
    assert lk.session.headers["Authorization"] == "Bearer key"
    res = lk.lookup("970436", "0071000123456")
    assert res.ok and res.account_name == "NGUYEN VAN A"
    assert sent == {"bank_code": "VCB", "bank_number": "0071000123456"}
    bad = lk.lookup("970422", "1")
    assert not bad.ok and "Không tìm thấy" in bad.message and sent["bank_code"] == "MB"


def test_fatal_error_stops_batch(matcher):
    from bank_checker.lookup import FatalLookupError

    calls = []

    def broke(b, a):
        calls.append(a)
        raise FatalLookupError("Số dư không đủ")

    df = pd.DataFrame({"STK": ["11111111", "22222222", "33333333"], "NH": ["VCB"] * 3})
    out = process(df, "STK", "NH", broke, matcher, delay=0)
    assert calls == ["11111111"]
    assert out.attrs["fatal_error"] == "Số dư không đủ"
    assert all("Số dư không đủ" in s for s in out["Trạng thái"])
